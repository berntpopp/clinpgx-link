"""Bounded OOXML admission before workbook object materialization."""

from __future__ import annotations

import io
import zipfile

import pytest
from openpyxl import Workbook


def _workbook_bytes() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["ID", "Value"])
    worksheet.append(["PA1", "normal"])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _replace_part(raw: bytes, path: str, transform) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(
        output, "w", compression=zipfile.ZIP_DEFLATED
    ) as target:
        for info in source.infolist():
            body = source.read(info)
            target.writestr(info, transform(body) if info.filename == path else body)
    return output.getvalue()


def _must_not_open(*_args, **_kwargs):
    raise AssertionError("unsafe OOXML reached openpyxl")


def _utf16_xml(body: bytes, transform=lambda value: value) -> bytes:
    text = transform(body.decode("utf-8"))
    return ('<?xml version="1.0" encoding="UTF-16"?>' + text).encode("utf-16")


def test_nested_expansion_limit_is_checked_before_openpyxl(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Catch inner ZIP expansion bypassing the outer archive/member limits."""
    from clinpgx_link.exceptions import DataValidationError
    from clinpgx_link.ingest import spreadsheets

    monkeypatch.setattr(spreadsheets, "load_workbook", _must_not_open)
    limits = spreadsheets.SpreadsheetLimits.for_tests(max_expanded_bytes=100)
    with pytest.raises(DataValidationError, match="expanded"):
        spreadsheets.parse_spreadsheet(io.BytesIO(_workbook_bytes()), limits=limits)


def test_xml_doctype_is_rejected_before_openpyxl(monkeypatch: pytest.MonkeyPatch) -> None:
    """Catch entity-bearing worksheet XML reaching a general OOXML parser."""
    from clinpgx_link.exceptions import DataValidationError
    from clinpgx_link.ingest import spreadsheets

    raw = _replace_part(
        _workbook_bytes(),
        "xl/worksheets/sheet1.xml",
        lambda body: body.replace(
            b"<worksheet",
            b'<!DOCTYPE worksheet [<!ENTITY x "boom">]><worksheet',
            1,
        ),
    )
    monkeypatch.setattr(spreadsheets, "load_workbook", _must_not_open)
    with pytest.raises(DataValidationError, match="declaration"):
        spreadsheets.parse_spreadsheet(io.BytesIO(raw))


def test_sparse_maximum_dimension_is_rejected_before_openpyxl(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Catch a tiny worksheet claiming an Excel-maximum materialized grid."""
    from clinpgx_link.exceptions import DataValidationError
    from clinpgx_link.ingest import spreadsheets

    raw = _replace_part(
        _workbook_bytes(),
        "xl/worksheets/sheet1.xml",
        lambda body: body.replace(b'ref="A1:B2"', b'ref="A1:XFD1048576"', 1),
    )
    monkeypatch.setattr(spreadsheets, "load_workbook", _must_not_open)
    with pytest.raises(DataValidationError, match="dimension"):
        spreadsheets.parse_spreadsheet(io.BytesIO(raw))


def test_giant_merge_area_is_rejected_before_openpyxl(monkeypatch: pytest.MonkeyPatch) -> None:
    """Catch merged-coordinate expansion consuming memory before row limits apply."""
    from clinpgx_link.exceptions import DataValidationError
    from clinpgx_link.ingest import spreadsheets

    def add_merge(body: bytes) -> bytes:
        return body.replace(
            b"</worksheet>",
            b'<mergeCells count="1"><mergeCell ref="A1:XFD1048576"/></mergeCells></worksheet>',
        )

    raw = _replace_part(_workbook_bytes(), "xl/worksheets/sheet1.xml", add_merge)
    monkeypatch.setattr(spreadsheets, "load_workbook", _must_not_open)
    with pytest.raises(DataValidationError, match="merge"):
        spreadsheets.parse_spreadsheet(io.BytesIO(raw))


def test_utf16_worksheet_cannot_bypass_dimension_admission(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Catch alternate XML encoding hiding a grid larger than the materialization cap."""
    from clinpgx_link.exceptions import DataValidationError
    from clinpgx_link.ingest import spreadsheets

    raw = _replace_part(
        _workbook_bytes(), "xl/worksheets/sheet1.xml", lambda body: _utf16_xml(body)
    )
    monkeypatch.setattr(spreadsheets, "load_workbook", _must_not_open)
    with pytest.raises(DataValidationError):
        spreadsheets.parse_spreadsheet(
            io.BytesIO(raw), limits=spreadsheets.SpreadsheetLimits.for_tests(max_cells=1)
        )


def test_utf16_worksheet_cannot_hide_giant_merge(monkeypatch: pytest.MonkeyPatch) -> None:
    """Catch alternate XML encoding hiding merged-coordinate expansion."""
    from clinpgx_link.exceptions import DataValidationError
    from clinpgx_link.ingest import spreadsheets

    def encoded(body: bytes) -> bytes:
        return _utf16_xml(
            body,
            lambda text: text.replace(
                "</worksheet>",
                '<mergeCells count="1"><mergeCell ref="A1:XFD1048576"/></mergeCells>'
                "</worksheet>",
            ),
        )

    raw = _replace_part(_workbook_bytes(), "xl/worksheets/sheet1.xml", encoded)
    monkeypatch.setattr(spreadsheets, "load_workbook", _must_not_open)
    with pytest.raises(DataValidationError):
        spreadsheets.parse_spreadsheet(io.BytesIO(raw))


def test_utf16_worksheet_cannot_hide_doctype(monkeypatch: pytest.MonkeyPatch) -> None:
    """Catch alternate XML encoding hiding an entity declaration from admission."""
    from clinpgx_link.exceptions import DataValidationError
    from clinpgx_link.ingest import spreadsheets

    raw = _replace_part(
        _workbook_bytes(),
        "xl/worksheets/sheet1.xml",
        lambda body: _utf16_xml(
            body,
            lambda text: '<!DOCTYPE worksheet [<!ENTITY x "boom">]>' + text,
        ),
    )
    monkeypatch.setattr(spreadsheets, "load_workbook", _must_not_open)
    with pytest.raises(DataValidationError):
        spreadsheets.parse_spreadsheet(io.BytesIO(raw))


def test_xml_namespace_variants_cannot_hide_sparse_cells(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Catch valid non-ASCII namespace prefixes that defeat byte-regex matching."""
    from clinpgx_link.exceptions import DataValidationError
    from clinpgx_link.ingest import spreadsheets

    raw = _replace_part(
        _workbook_bytes(),
        "xl/worksheets/sheet1.xml",
        lambda body: body.replace(
            b'<c r="A1"',
            '<é:c xmlns:é="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'r="C1"'.encode(),
            1,
        ).replace(b"</c>", "</é:c>".encode(), 1),
    )
    monkeypatch.setattr(spreadsheets, "load_workbook", _must_not_open)
    with pytest.raises(DataValidationError, match="cell count or coordinate"):
        spreadsheets.parse_spreadsheet(
            io.BytesIO(raw), limits=spreadsheets.SpreadsheetLimits.for_tests(max_columns=2)
        )
