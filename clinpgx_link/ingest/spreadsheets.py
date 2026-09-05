"""Profiled OOXML workbook reader preserving cell-level source semantics."""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass
from typing import Any, BinaryIO

from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

from clinpgx_link.exceptions import DataValidationError


@dataclass(frozen=True)
class SpreadsheetCell:
    value: Any
    formula: str | None
    cached_value: Any
    present: bool
    merged_range: str | None
    merge_anchor: str | None


@dataclass(frozen=True)
class SpreadsheetRow:
    ordinal: int
    fields: dict[str, SpreadsheetCell]


@dataclass(frozen=True)
class SpreadsheetSheet:
    name: str
    headers: tuple[str, ...]
    rows: tuple[SpreadsheetRow, ...]
    hidden: bool
    merged_ranges: tuple[str, ...]


@dataclass(frozen=True)
class SpreadsheetDocument:
    sheets: tuple[SpreadsheetSheet, ...]


def _workbook_bytes(stream: BinaryIO) -> bytes:
    chunks: list[bytes] = []
    while chunk := stream.read(1024 * 1024):
        chunks.append(chunk)
    raw = b"".join(chunks)
    if not zipfile.is_zipfile(io.BytesIO(raw)):
        raise DataValidationError("Spreadsheet does not have an OOXML ZIP signature")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as package:
            members = set(package.namelist())
    except zipfile.BadZipFile as exc:
        raise DataValidationError("Spreadsheet OOXML package is invalid") from exc
    if "[Content_Types].xml" not in members or "xl/workbook.xml" not in members:
        raise DataValidationError("Spreadsheet is missing required OOXML parts")
    return raw


def _merged_cells(worksheet: Any) -> dict[str, tuple[str, str]]:
    result: dict[str, tuple[str, str]] = {}
    for merged in worksheet.merged_cells.ranges:
        range_name = str(merged)
        anchor = f"{get_column_letter(merged.min_col)}{merged.min_row}"
        for row in range(merged.min_row, merged.max_row + 1):
            for column in range(merged.min_col, merged.max_col + 1):
                result[f"{get_column_letter(column)}{row}"] = (range_name, anchor)
    return result


def parse_spreadsheet(stream: BinaryIO) -> SpreadsheetDocument:
    """Read a profiled first-row-header workbook without flattening formulas/merges."""
    raw = _workbook_bytes(stream)
    try:
        formulas = load_workbook(io.BytesIO(raw), data_only=False, read_only=False)
        cached = load_workbook(io.BytesIO(raw), data_only=True, read_only=False)
    except (KeyError, OSError, ValueError, zipfile.BadZipFile) as exc:
        raise DataValidationError("Spreadsheet OOXML parts cannot be decoded") from exc
    sheets: list[SpreadsheetSheet] = []
    try:
        for worksheet in formulas.worksheets:
            cached_sheet = cached[worksheet.title]
            max_column = worksheet.max_column
            if max_column < 1 or worksheet.max_row < 1:
                raise DataValidationError("Spreadsheet sheet is empty")
            headers: list[str] = []
            for column in range(1, max_column + 1):
                value = worksheet.cell(row=1, column=column).value
                if not isinstance(value, str) or not value:
                    raise DataValidationError("Spreadsheet first row must contain text headers")
                headers.append(value)
            if len(headers) != len(set(headers)):
                raise DataValidationError("Spreadsheet has duplicate headers")
            existing = getattr(worksheet, "_cells", {})
            present = {
                f"{get_column_letter(column)}{row}"
                for row, column in existing
                if existing[(row, column)].__class__.__name__ != "MergedCell"
            }
            merges = _merged_cells(worksheet)
            rows: list[SpreadsheetRow] = []
            for row_number in range(2, worksheet.max_row + 1):
                fields: dict[str, SpreadsheetCell] = {}
                for column, header in enumerate(headers, start=1):
                    coordinate = f"{get_column_letter(column)}{row_number}"
                    cell = worksheet[coordinate]
                    formula = (
                        cell.value
                        if cell.data_type == "f" and isinstance(cell.value, str)
                        else None
                    )
                    cached_value = cached_sheet[coordinate].value if formula is not None else None
                    merge = merges.get(coordinate)
                    fields[header] = SpreadsheetCell(
                        value=cached_value if formula is not None else cell.value,
                        formula=formula,
                        cached_value=cached_value,
                        present=coordinate in present,
                        merged_range=merge[0] if merge else None,
                        merge_anchor=merge[1] if merge else None,
                    )
                rows.append(SpreadsheetRow(ordinal=row_number - 1, fields=fields))
            sheets.append(
                SpreadsheetSheet(
                    name=worksheet.title,
                    headers=tuple(headers),
                    rows=tuple(rows),
                    hidden=worksheet.sheet_state != "visible",
                    merged_ranges=tuple(str(item) for item in worksheet.merged_cells.ranges),
                )
            )
    finally:
        formulas.close()
        cached.close()
    return SpreadsheetDocument(sheets=tuple(sheets))


__all__ = [
    "SpreadsheetCell",
    "SpreadsheetDocument",
    "SpreadsheetRow",
    "SpreadsheetSheet",
    "parse_spreadsheet",
]
